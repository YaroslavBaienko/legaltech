#!/usr/bin/env python3
"""Exhaustive Cross-Validation & Cross-Verification Test Suite for Marma-Crypto Protocol v1.2.0.

Covers:
1. Canonical 9-Triangle Sri Yantra Geometry & Bilateral Symmetry Invariants (lx + rx == 300.0).
2. Cross-Runtime Bit-Exactness: Python (`marma_watermark.py`) <-> Node.js (`dstu4145_bridge.js`) <-> Browser JS (`docs/index.html`).
3. Pure-JS FIPS 180-4 SHA-256 & Python-compatible `computePublicSealDigest` cross-verification.
4. DSTU 4145-2002 Elliptic Curve GF(2^257), GOST 34.311-95, DSTU 7564:2014 (Kupyna-256), and RFC 6979 determinism.
5. Full 36-Coordinate Tamper Matrix (lx, rx, by, ay across all 9 triangles) in both Python and Browser JS.
6. Multi-format (PDF modes both/watermark/seal, DOCX with tables, SVG, and Self-Sealed Protocol PDF) round-trip verification.
7. Web Verifier (`docs/index.html`) Attack Simulation & Reset State Machine + Language Toggle + Same-File Re-Upload.
8. Production Whitepaper & PAdES-LTA / CAdES Artifact & Manifest Integrity Audit.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import docx
import pymupdf

PACKAGE_ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = PACKAGE_ROOT / "src"
REPO_ROOT = PACKAGE_ROOT
if not (REPO_ROOT / "docs" / "index.html").exists() and (PACKAGE_ROOT / "subproject_link.json").exists():
    _link = json.loads((PACKAGE_ROOT / "subproject_link.json").read_text(encoding="utf-8"))
    if _link.get("canonical_source_dir") and Path(_link["canonical_source_dir"]).exists():
        REPO_ROOT = Path(_link["canonical_source_dir"])
sys.path.insert(0, str(SRC_DIR))

import marma_watermark as mw  # noqa: E402


def run_browser_js_harness(js_body: str) -> dict:
    """Execute a test snippet inside the actual `docs/index.html` JavaScript VM context via Node.js."""
    script_code = f"""
const fs = require("fs");
const vm = require("vm");
const html = fs.readFileSync({json.dumps(str(REPO_ROOT / "docs" / "index.html"))}, "utf8");
const scriptMatch = html.match(/<script>([\\s\\S]*?)<\\/script>/);
if (!scriptMatch) throw new Error("No <script> tag in docs/index.html");

const elems = {{}};
const sandbox = {{
  document: {{
    documentElement: {{ lang: "en" }},
    getElementById: (id) => {{
      if (!elems[id]) {{
        elems[id] = {{
          id,
          textContent: "",
          innerHTML: "",
          value: "",
          style: {{}},
          className: "",
          classList: {{ add(){{}}, remove(){{}} }},
          addEventListener(){{}},
          appendChild(){{}}
        }};
      }}
      return elems[id];
    }},
    querySelectorAll: () => [],
    createElement: () => ({{ innerHTML: "", className: "", style: {{}} }})
  }},
  TextEncoder,
  TextDecoder,
  DataView,
  Uint8Array,
  Uint32Array,
  Buffer,
  Response,
  DecompressionStream: globalThis.DecompressionStream,
  console,
  elems,
  fs
}};
vm.createContext(sandbox);
vm.runInContext(scriptMatch[1], sandbox);
const testFn = new vm.Script({json.dumps(f"(async () => {{ {js_body} }})()")});
testFn.runInContext(sandbox).then(res => {{
  console.log(JSON.stringify({{ ok: true, result: res }}));
}}).catch(err => {{
  console.log(JSON.stringify({{ ok: false, error: err.message || String(err) }}));
  process.exit(1);
}});
"""
    proc = subprocess.run(
        ["node", "-e", script_code],
        capture_output=True,
        text=True,
        timeout=25,
        check=False,
    )
    if proc.returncode != 0:
        raise AssertionError(f"Browser JS harness failed: {proc.stderr or proc.stdout}")
    last_line = proc.stdout.strip().splitlines()[-1]
    payload = json.loads(last_line)
    if not payload.get("ok"):
        raise AssertionError(f"Browser JS error: {payload.get('error')}")
    return payload["result"]


class TestMarmaCryptoCrossVerification(unittest.TestCase):
    """Comprehensive audit, cross-validation, and cross-verification test suite."""

    def test_01_canonical_geometry_and_bilateral_symmetry(self) -> None:
        """Verify all 9 canonical triangles satisfy lx + rx == 300.0 and valid ordering."""
        self.assertEqual(len(mw.CANONICAL_TRIANGLES), 9)
        names = [t[0] for t in mw.CANONICAL_TRIANGLES]
        self.assertEqual(names, ["D1", "U1", "U3", "U2", "D3", "D2", "U4", "D4", "D5"])
        for name, lx, by, ay, rx, is_up in mw.CANONICAL_TRIANGLES:
            self.assertAlmostEqual(lx + rx, 300.0, places=9, msg=f"Symmetry broken on {name}")
            self.assertGreater(rx, lx, msg=f"rx <= lx on {name}")
            self.assertEqual(name.startswith("U"), is_up)
            if is_up:
                self.assertGreater(by, ay, msg=f"Upward triangle {name} must have by > ay")
            else:
                self.assertLess(by, ay, msg=f"Downward triangle {name} must have by < ay")

    def test_02_cross_runtime_sha256_and_perturbation_derivation(self) -> None:
        """Cross-verify Python hashlib/perturbation math against Browser JS sha256Hex/applySignedPerturbations."""
        test_strings = [
            "",
            "abc",
            "MARMA-CRYPTO-PROTOCOL-V1.2",
            "Баєнко Ярослав Володимирович :: ДСТУ 4145-2002 :: Шрі Янтра 54 марми",
            "A" * 256,
        ]
        expected_hashes = [hashlib.sha256(s.encode("utf-8")).hexdigest() for s in test_strings]

        # Generate 20 deterministic HMAC hex digests and compare Python vs JS offsets & triangle coordinates
        sample_hmacs = [
            hashlib.sha256(f"SEED-{i}".encode("utf-8")).hexdigest() for i in range(20)
        ]
        py_offsets = [mw.derive_offsets_from_hmac_hex(h) for h in sample_hmacs]

        js_res = run_browser_js_harness(f"""
          const strings = {json.dumps(test_strings)};
          const jsHashes = strings.map(s => sha256Hex(s));
          const hmacs = {json.dumps(sample_hmacs)};
          const jsOffsets = hmacs.map(h => deriveOffsetsFromHmac(h));
          const tris150 = applySignedPerturbations(solveBaselineYantra(150.0), hmacs[0], 150.0);
          const tris300 = applySignedPerturbations(solveBaselineYantra(300.0), hmacs[0], 300.0);
          return {{ jsHashes, jsOffsets, tris150, tris300 }};
        """)

        self.assertEqual(js_res["jsHashes"], expected_hashes)
        self.assertEqual(js_res["jsOffsets"], py_offsets)

        # Verify coordinate scaling [0, 300]^2 and [0, 600]^2 between Python and JS
        d150, s150 = mw.verify_triangle_coordinates(js_res["tris150"], py_offsets[0], cx=150.0)
        d300, s300 = mw.verify_triangle_coordinates(js_res["tris300"], py_offsets[0], cx=300.0)
        self.assertEqual(d150, 0.0)
        self.assertEqual(s150, 0.0)
        self.assertEqual(d300, 0.0)
        self.assertEqual(s300, 0.0)

    def test_03_dstu4145_bridge_sign_verify_and_tamper_rejection(self) -> None:
        """Cross-verify dstu4145_bridge.js RFC 6979 determinism, GF(2^257) EC math, and tamper rejection."""
        payload_hex = "EDU-VEDI::CROSS-VERIFICATION-TEST::2026".encode("utf-8").hex()
        sign1 = mw.call_dstu4145_bridge({
            "action": "sign",
            "secretSeed": "Test-Key-Seed-4145",
            "payloadBytesHex": payload_hex,
        })
        sign2 = mw.call_dstu4145_bridge({
            "action": "sign",
            "secretSeed": "Test-Key-Seed-4145",
            "payloadBytesHex": payload_hex,
        })
        self.assertEqual(sign1["status"], "ok")
        # RFC 6979 determinism: signing identical payload with identical key must yield identical (r, s)
        self.assertEqual(sign1["r_hex"], sign2["r_hex"])
        self.assertEqual(sign1["s_hex"], sign2["s_hex"])
        self.assertEqual(sign1["rs_digest_hex"], sign2["rs_digest_hex"])
        self.assertEqual(
            sign1["rs_digest_hex"],
            hashlib.sha256(f"{sign1['r_hex']}::{sign1['s_hex']}".encode("utf-8")).hexdigest(),
        )

        # Valid verification
        ver_ok = mw.call_dstu4145_bridge({
            "action": "verify",
            "payloadBytesHex": payload_hex,
            "dstu4145": sign1,
        })
        self.assertTrue(ver_ok["dstu4145_valid"])
        self.assertTrue(ver_ok["gost34311_match"])
        self.assertTrue(ver_ok["ec_signature_valid"])
        self.assertTrue(ver_ok["rs_digest_match"])

        # Tamper 1: Modified payload bytes
        ver_bad_payload = mw.call_dstu4145_bridge({
            "action": "verify",
            "payloadBytesHex": "EDU-VEDI::CROSS-VERIFICATION-TAMPERED".encode("utf-8").hex(),
            "dstu4145": sign1,
        })
        self.assertFalse(ver_bad_payload["dstu4145_valid"])
        self.assertFalse(ver_bad_payload["gost34311_match"])

        # Tamper 2: Modified vector r_hex on GF(2^257)
        bad_r = dict(sign1)
        bad_r["r_hex"] = sign1["r_hex"][:-2] + ("00" if sign1["r_hex"][-2:] != "00" else "ff")
        ver_bad_r = mw.call_dstu4145_bridge({
            "action": "verify",
            "payloadBytesHex": payload_hex,
            "dstu4145": bad_r,
        })
        self.assertFalse(ver_bad_r["dstu4145_valid"])

        # Tamper 3: Modified rs_digest_hex
        bad_rs_digest = dict(sign1)
        bad_rs_digest["rs_digest_hex"] = "0" * 64
        ver_bad_rs = mw.call_dstu4145_bridge({
            "action": "verify",
            "payloadBytesHex": payload_hex,
            "dstu4145": bad_rs_digest,
        })
        self.assertFalse(ver_bad_rs["dstu4145_valid"])
        self.assertFalse(ver_bad_rs["rs_digest_match"])

    def test_04_36_coordinate_tamper_matrix_python_and_js(self) -> None:
        """Perturb every single coordinate (lx, rx, by, ay) across all 9 triangles (36 tests) and verify 100% detection."""
        with tempfile.TemporaryDirectory(prefix="marma_36coord_") as tmp:
            svg_path = Path(tmp) / "matrix_test.svg"
            payload = "MARMA-36-COORDINATE-AUDIT-MATRIX"
            mw.sign_sri_yantra_svg(payload, mw.DEFAULT_SECRET, svg_path)
            base_svg = svg_path.read_text(encoding="utf-8")

            # Clean baseline must pass in both Python and Browser JS
            py_clean = mw.verify_sri_yantra_svg(base_svg, payload)
            self.assertTrue(py_clean["verified"])
            self.assertEqual(py_clean["max_coordinate_signature_drift"], 0.0)
            self.assertEqual(py_clean["max_bilateral_symmetry_error"], 0.0)

            for tri in py_clean["triangles"]:
                t_name = tri["name"]
                for coord_key in ("lx", "by", "ay", "rx"):
                    orig_val = f'{coord_key}="{tri[coord_key]:.9f}"'
                    tampered_val = f'{coord_key}="{tri[coord_key] + 0.000100000:.9f}"'
                    old_tag_prefix = f'id="tri_{t_name}"'
                    # Replace only inside the target triangle tag
                    lines = []
                    for line in base_svg.splitlines():
                        if old_tag_prefix in line:
                            self.assertIn(orig_val, line)
                            line = line.replace(orig_val, tampered_val, 1)
                        lines.append(line)
                    tampered_svg = "\n".join(lines)

                    py_tamper = mw.verify_sri_yantra_svg(tampered_svg, payload)
                    self.assertFalse(
                        py_tamper["verified"],
                        msg=f"Python verifier missed tamper on {t_name}.{coord_key}",
                    )
                    self.assertGreaterEqual(
                        py_tamper["max_coordinate_signature_drift"],
                        9e-5,
                        msg=f"Python drift not detected on {t_name}.{coord_key}",
                    )

            # Also verify in Browser JS that tampering `by`, `ay`, `public_seal_digest`, or `r_hex` alone is caught!
            js_check = run_browser_js_harness(f"""
              const buf = fs.readFileSync({json.dumps(str(svg_path))});
              const seal = extractMarmaSealFromBytes(new Uint8Array(buf), "matrix_test.svg");
              renderVerificationState(seal, "matrix_test.svg", false);
              const cleanOk = elems.statusBanner.className.includes("ok");

              // Tamper ONLY ay of D5 by +0.0001 (lx + rx remains 300.0!)
              const sealAyTamper = JSON.parse(JSON.stringify(seal));
              sealAyTamper.triangles[8].ay = Number((sealAyTamper.triangles[8].ay + 0.0001).toFixed(9));
              renderVerificationState(sealAyTamper, "ay_tamper.svg", true);
              const ayTamperRejected = elems.statusBanner.className.includes("err");

              // Tamper ONLY by of U4 by -0.0001 (lx + rx remains 300.0!)
              const sealByTamper = JSON.parse(JSON.stringify(seal));
              sealByTamper.triangles[6].by = Number((sealByTamper.triangles[6].by - 0.0001).toFixed(9));
              renderVerificationState(sealByTamper, "by_tamper.svg", true);
              const byTamperRejected = elems.statusBanner.className.includes("err");

              return {{ cleanOk, ayTamperRejected, byTamperRejected }};
            """)
            self.assertTrue(js_check["cleanOk"])
            self.assertTrue(js_check["ayTamperRejected"])
            self.assertTrue(js_check["byTamperRejected"])

    def test_05_pdf_docx_svg_and_protocol_end_to_end_cross_runtime(self) -> None:
        """Sign PDF (all 3 modes), DOCX (paragraphs + tables), SVG, and Protocol PDF; verify in Python & Browser JS."""
        with tempfile.TemporaryDirectory(prefix="marma_e2e_") as tmp:
            tmp_dir = Path(tmp)

            # 1. PDF across modes: both, watermark, seal
            raw_pdf = tmp_dir / "source.pdf"
            doc = pymupdf.open()
            for p_idx in range(2):
                page = doc.new_page()
                page.insert_text((72, 100), f"Marma-Crypto Test Page {p_idx + 1} — Правовий протокол Шрі Янтри", fontsize=12)
            doc.save(str(raw_pdf))
            doc.close()

            for mode in ("both", "watermark", "seal"):
                signed_pdf = tmp_dir / f"signed_{mode}.pdf"
                mw.stamp_pdf_document(raw_pdf, signed_pdf, secret_key=mw.DEFAULT_SECRET, mode=mode)
                py_res = mw.verify_pdf_document(signed_pdf)
                self.assertTrue(py_res["verified"], f"PDF mode={mode} failed Python verification")
                self.assertEqual(py_res["max_coordinate_signature_drift"], 0.0)
                self.assertEqual(py_res["max_bilateral_symmetry_error"], 0.0)

                js_res = run_browser_js_harness(f"""
                  const buf = fs.readFileSync({json.dumps(str(signed_pdf))});
                  const seal = extractMarmaSealFromBytes(new Uint8Array(buf), {json.dumps(signed_pdf.name)});
                  renderVerificationState(seal, {json.dumps(signed_pdf.name)}, false);
                  return {{
                    banner: elems.statusBanner.className,
                    errors: elems.valErrors.textContent
                  }};
                """)
                self.assertIn("ok", js_res["banner"])
                self.assertIn("Δ_sym = 0.000000000000  |  Δ_coord = 0.000000000000", js_res["errors"])

            # 2. Protocol PDF generation & self-verification in Python and Browser JS
            proto_pdf = tmp_dir / "signed_both_protocol.pdf"
            mw.generate_pdf_verification_protocol(tmp_dir / "signed_both.pdf", proto_pdf, lang="uk")
            proto_py = mw.verify_pdf_document(proto_pdf)
            self.assertTrue(proto_py["verified"], "Generated Protocol PDF must itself be verifiable")

            proto_js = run_browser_js_harness(f"""
              const buf = fs.readFileSync({json.dumps(str(proto_pdf))});
              const seal = extractMarmaSealFromBytes(new Uint8Array(buf), "signed_both_protocol.pdf");
              renderVerificationState(seal, "signed_both_protocol.pdf", false);
              return {{ banner: elems.statusBanner.className, errors: elems.valErrors.textContent }};
            """)
            self.assertIn("ok", proto_js["banner"])

            # 3. DOCX with paragraphs and tables — verify in Python and Browser JS (via DecompressionStream)
            raw_docx = tmp_dir / "contract.docx"
            signed_docx = tmp_dir / "contract_signed.docx"
            w = docx.Document()
            w.add_heading("Договір про передачу прав інтелектуальної власності", 0)
            w.add_paragraph("Сторона 1: Баєнко Ярослав Володимирович.")
            tbl = w.add_table(rows=2, cols=2)
            tbl.cell(0, 0).text = "Параметр"
            tbl.cell(0, 1).text = "Значення"
            tbl.cell(1, 0).text = "Стандарт"
            tbl.cell(1, 1).text = "ДСТУ 4145-2002 / Шрі Янтра"
            w.save(str(raw_docx))

            mw.stamp_docx_document(raw_docx, signed_docx, secret_key=mw.DEFAULT_SECRET)
            docx_py = mw.verify_docx_document(signed_docx)
            self.assertTrue(docx_py["verified"])
            self.assertEqual(docx_py["max_coordinate_signature_drift"], 0.0)

            docx_js = run_browser_js_harness(f"""
              const buf = fs.readFileSync({json.dumps(str(signed_docx))});
              const seal = await extractDocxMarmaSeal(new Uint8Array(buf));
              renderVerificationState(seal, "contract_signed.docx", false);
              return {{ banner: elems.statusBanner.className, errors: elems.valErrors.textContent }};
            """)
            self.assertIn("ok", docx_js["banner"])
            self.assertIn("Δ_sym = 0.000000000000  |  Δ_coord = 0.000000000000", docx_js["errors"])

            # Tamper DOCX table cell text and verify Python rejects it
            tampered_docx = tmp_dir / "contract_tampered.docx"
            w_mod = docx.Document(str(signed_docx))
            w_mod.tables[0].cell(1, 1).text = "ДСТУ 4145-2002 (ПІДРОБЛЕНО)"
            w_mod.save(str(tampered_docx))
            docx_tamper_py = mw.verify_docx_document(tampered_docx)
            self.assertFalse(docx_tamper_py["verified"])
            self.assertIn("modified", docx_tamper_py["reason"])

    def test_06_web_verifier_attack_simulation_and_reset_lifecycle(self) -> None:
        """Exhaustively test simulateTamperAttack(), toggle reset, language switch, and re-uploading after attack."""
        res = run_browser_js_harness(f"""
          const states = [];
          // 1. Initial load
          states.push({{
            step: "init",
            banner: elems.statusBanner.className,
            btn: elems.btnSimulateAttack.textContent,
            errors: elems.valErrors.textContent
          }});

          // 2. Click Simulate Attack
          simulateTamperAttack();
          states.push({{
            step: "attack_on",
            banner: elems.statusBanner.className,
            btn: elems.btnSimulateAttack.textContent,
            errors: elems.valErrors.textContent
          }});

          // 3. Language toggle while attack is active (starts in EN -> switches to UK -> back to EN)
          const btnEnAttack = elems.btnSimulateAttack.textContent;
          toggleLanguage();
          const btnUkAttack = elems.btnSimulateAttack.textContent;
          toggleLanguage();

          // 4. Click Reset Attack (second click on same button)
          simulateTamperAttack();
          states.push({{
            step: "attack_reset",
            banner: elems.statusBanner.className,
            btn: elems.btnSimulateAttack.textContent,
            errors: elems.valErrors.textContent
          }});

          // 5. Activate attack again, then load each of the 5 repository artifacts
          const artifactFiles = [
            "docs/whitepaper/Marma_Crypto_Whitepaper_UA.pdf",
            "docs/whitepaper/Marma_Crypto_Whitepaper_EN.pdf",
            "docs/whitepaper/sri_yantra_whitepaper_diagram.svg",
            "docs/whitepaper/Marma_Crypto_Whitepaper_UA_PROTOCOL.pdf",
            "docs/whitepaper/Marma_Crypto_Whitepaper_EN_PROTOCOL.pdf"
          ];
          const postAttackLoads = [];
          for (const relPath of artifactFiles) {{
            simulateTamperAttack(); // put into tampered state first!
            if (!elems.statusBanner.className.includes("err")) {{
              throw new Error("Expected err state before loading " + relPath);
            }}
            const fullPath = {json.dumps(str(REPO_ROOT))} + "/" + relPath;
            const buf = fs.readFileSync(fullPath);
            const seal = extractMarmaSealFromBytes(new Uint8Array(buf), relPath);
            renderVerificationState(seal, relPath, false);
            postAttackLoads.push({{
              file: relPath,
              banner: elems.statusBanner.className,
              btn: elems.btnSimulateAttack.textContent,
              errors: elems.valErrors.textContent,
              tsp: elems.valStandard ? elems.valStandard.textContent : ""
            }});
          }}

          return {{ states, btnEnAttack, btnUkAttack, postAttackLoads, defaultLang: currentLang }};
        """)

        self.assertEqual(res["defaultLang"], "en")
        self.assertEqual(res["states"][0]["banner"], "status-banner ok")
        self.assertEqual(res["states"][1]["banner"], "status-banner err")
        self.assertIn("Reset Attack", res["states"][1]["btn"])
        self.assertIn("Reset Attack", res["btnEnAttack"])
        self.assertIn("Скинути атаку", res["btnUkAttack"])
        self.assertEqual(res["states"][2]["banner"], "status-banner ok")
        self.assertEqual(
            res["states"][2]["errors"],
            "Δ_sym = 0.000000000000  |  Δ_coord = 0.000000000000",
        )

        for item in res["postAttackLoads"]:
            self.assertEqual(item["banner"], "status-banner ok", f"Failed post-attack load on {item['file']}")
            self.assertEqual(
                item["errors"],
                "Δ_sym = 0.000000000000  |  Δ_coord = 0.000000000000",
                f"Non-zero error on {item['file']}",
            )
            self.assertIn("Simulate Attack", item["btn"])
            self.assertIn("2026-10-17 03:20:22 UTC", item["tsp"])

    def test_07_production_whitepapers_pades_lta_and_manifests_audit(self) -> None:
        """Cross-audit production Whitepapers, PAdES-LTA containers, SHA256SUMS.txt, and directory sync."""
        wp_dir = REPO_ROOT / "whitepaper"
        docs_wp_dir = REPO_ROOT / "docs" / "whitepaper"

        # 1. Verify SHA256SUMS.txt matches actual files
        sha_lines = (wp_dir / "SHA256SUMS.txt").read_text(encoding="utf-8").strip().splitlines()
        self.assertGreaterEqual(len(sha_lines), 6)
        for line in sha_lines:
            expected_sha, fname = line.split()
            actual_sha = hashlib.sha256((wp_dir / fname).read_bytes()).hexdigest()
            self.assertEqual(expected_sha, actual_sha, f"SHA256SUMS mismatch for {fname}")

        # 2. Verify docs/whitepaper copies are byte-for-byte identical to whitepaper/
        for shared_name in [
            "Marma_Crypto_Whitepaper_UA.pdf",
            "Marma_Crypto_Whitepaper_EN.pdf",
            "Marma_Crypto_Whitepaper_UA_PROTOCOL.pdf",
            "Marma_Crypto_Whitepaper_EN_PROTOCOL.pdf",
            "sri_yantra_whitepaper_diagram.svg",
            "sri_yantra_whitepaper_diagram.svg.p7s",
            "AUTHORSHIP_MANIFEST.txt",
        ]:
            h1 = hashlib.sha256((wp_dir / shared_name).read_bytes()).hexdigest()
            h2 = hashlib.sha256((docs_wp_dir / shared_name).read_bytes()).hexdigest()
            self.assertEqual(h1, h2, f"Drift between whitepaper/{shared_name} and docs/whitepaper/{shared_name}")

        # 3. Verify PAdES-LTA and DSTU 4145-2002 EC signatures on both Whitepaper PDFs
        for pdf_name in ("Marma_Crypto_Whitepaper_UA.pdf", "Marma_Crypto_Whitepaper_EN.pdf"):
            res = mw.verify_pdf_document(wp_dir / pdf_name)
            self.assertTrue(res["verified"], f"{pdf_name} failed verification")
            self.assertTrue(res["dstu4145_verified"])
            self.assertTrue(res["dstu4145_ec_verified"])
            self.assertTrue(res["pades_cades_verified"])
            self.assertTrue(res["hmac_match"])
            self.assertTrue(res["public_seal_match"])
            self.assertEqual(res["max_coordinate_signature_drift"], 0.0)
            self.assertEqual(res["max_bilateral_symmetry_error"], 0.0)
            self.assertEqual(res["signer_info"]["rnokpp"], "3255419579")
            qes_rep = res["qes_container_report"]
            self.assertEqual(qes_rep["signature_level"], "PAdES-LTA")
            self.assertTrue(qes_rep["has_tsp"])
            self.assertTrue(qes_rep["has_ltv"])
            self.assertTrue(qes_rep["has_ocsp"])
            self.assertTrue(qes_rep["has_pdf_dss"])
            self.assertTrue(qes_rep["is_valid_at_signing"])

        # 4. Verify verify_any_file resolves .typ -> .pdf and .svg.p7s -> .svg
        res_typ = mw.verify_any_file(wp_dir / "Marma_Crypto_Whitepaper_EN.typ")
        self.assertTrue(res_typ["verified"])
        res_p7s = mw.verify_any_file(wp_dir / "sri_yantra_whitepaper_diagram.svg.p7s")
        self.assertTrue(res_p7s["verified"])
        self.assertTrue(res_p7s["pades_cades_verified"])

    def test_08_cryptographic_commitment_and_edge_case_rejections(self) -> None:
        """Verify rejection of public_seal_digest tamper, r||s binding tamper, and unsigned/corrupted files."""
        with tempfile.TemporaryDirectory(prefix="marma_edge_") as tmp:
            tmp_dir = Path(tmp)
            svg_p = tmp_dir / "signed.svg"
            mw.sign_sri_yantra_svg("EDGE-CASE-PAYLOAD", mw.DEFAULT_SECRET, svg_p)
            svg_txt = svg_p.read_text(encoding="utf-8")

            # 1. Tamper data-public-seal in SVG while keeping all coordinates & HMAC intact
            bad_pseal_svg = svg_txt.replace('data-public-seal="', 'data-public-seal="0000', 1)
            bad_pseal_svg = bad_pseal_svg.replace('0000' + svg_txt.split('data-public-seal="')[1][:64], "0" * 64, 1)
            py_bad_pseal = mw.verify_sri_yantra_svg(bad_pseal_svg)
            self.assertFalse(py_bad_pseal["verified"])
            self.assertFalse(py_bad_pseal["public_seal_match"])

            # 2. Tamper r||s binding and public_seal_digest in Browser JS
            js_edge = run_browser_js_harness(f"""
              const buf = fs.readFileSync({json.dumps(str(svg_p))});
              const seal = extractMarmaSealFromBytes(new Uint8Array(buf), "signed.svg");

              // Tamper public_seal_digest
              const badSeal = JSON.parse(JSON.stringify(seal));
              badSeal.public_seal_digest = "f".repeat(64);
              renderVerificationState(badSeal, "bad_pseal.svg", true);
              const psealRejected = elems.statusBanner.className.includes("err");

              // Tamper dstu4145.r_hex (breaks SHA256(r||s) == hmac_sha256)
              const badRS = JSON.parse(JSON.stringify(seal));
              badRS.dstu4145.r_hex = "1".repeat(64);
              renderVerificationState(badRS, "bad_rs.svg", true);
              const rsRejected = elems.statusBanner.className.includes("err");

              // Tamper payload (breaks SHA256(payload) == content_sha256)
              const badPay = JSON.parse(JSON.stringify(seal));
              badPay.payload = "FORGED-PAYLOAD";
              renderVerificationState(badPay, "bad_pay.svg", true);
              const payRejected = elems.statusBanner.className.includes("err");

              return {{ psealRejected, rsRejected, payRejected }};
            """)
            self.assertTrue(js_edge["psealRejected"])
            self.assertTrue(js_edge["rsRejected"])
            self.assertTrue(js_edge["payRejected"])

            # 3. Unsigned PDF and DOCX must be cleanly rejected
            un_pdf = tmp_dir / "unsigned.pdf"
            d = pymupdf.open()
            d.new_page().insert_text((72, 72), "Unsigned PDF")
            d.save(str(un_pdf))
            d.close()
            self.assertFalse(mw.verify_pdf_document(un_pdf)["verified"])

            un_docx = tmp_dir / "unsigned.docx"
            w = docx.Document()
            w.add_paragraph("Unsigned DOCX")
            w.save(str(un_docx))
            self.assertFalse(mw.verify_docx_document(un_docx)["verified"])

    def test_09_official_ukrainian_iit_euscp_czo_verification_and_tsp_timestamp(self) -> None:
        """Directly verify Whitepapers & SVG with official Ukrainian IIT euscp.so + pdfs.so (czo.gov.ua/verify engine)."""
        iit_lib = Path("/opt/iit/eu/sw/euscp.so")
        if not iit_lib.exists():
            self.skipTest("/opt/iit/eu/sw/euscp.so not installed")

        code = f"""
import ctypes, json, sys
sys.path.insert(0, '/home/attor/legaltech/packages/qes-tools/src')
from token_bridge import load_euscp_lib, init_eusign

class SYSTEMTIME(ctypes.Structure):
    _fields_ = [
        ("wYear", ctypes.c_ushort),
        ("wMonth", ctypes.c_ushort),
        ("wDayOfWeek", ctypes.c_ushort),
        ("wDay", ctypes.c_ushort),
        ("wHour", ctypes.c_ushort),
        ("wMinute", ctypes.c_ushort),
        ("wSecond", ctypes.c_ushort),
        ("wMilliseconds", ctypes.c_ushort),
    ]

class EU_TIME_INFO(ctypes.Structure):
    _fields_ = [
        ("dwVersion", ctypes.c_ulong),
        ("bTimeAvail", ctypes.c_int),
        ("bTimeStamp", ctypes.c_int),
        ("Time", SYSTEMTIME),
        ("bSignTimeStampAvail", ctypes.c_int),
        ("SignTimeStamp", SYSTEMTIME),
    ]

lib = load_euscp_lib()
init_eusign(lib)
lib.EUSetUIMode(False)

out = {{}}
for name in ["Marma_Crypto_Whitepaper_EN.pdf", "Marma_Crypto_Whitepaper_UA.pdf"]:
    data = open({json.dumps(str(REPO_ROOT / "whitepaper"))} + "/" + name, "rb").read()
    buf = ctypes.create_string_buffer(data, len(data))
    cnt = ctypes.c_uint32(0)
    rv_cnt = lib.EUPDFGetSignsCount(buf, len(data), ctypes.byref(cnt))
    info_buf = ctypes.create_string_buffer(512)
    rv_ver = lib.EUPDFVerifyData(0, buf, len(data), info_buf)
    p_time = ctypes.POINTER(EU_TIME_INFO)()
    rv_time = lib.EUPDFGetSignTimeInfo(0, buf, len(data), ctypes.byref(p_time))
    st = ctypes.c_uint32(0)
    rv_type = lib.EUPDFGetSignType(0, buf, len(data), ctypes.byref(st))
    out[name] = {{
        "rv_cnt": rv_cnt,
        "count": cnt.value,
        "rv_ver": rv_ver,
        "rv_time": rv_time,
        "bTimeAvail": p_time.contents.bTimeAvail if rv_time == 0 else 0,
        "bTimeStamp": p_time.contents.bTimeStamp if rv_time == 0 else 0,
        "year": p_time.contents.Time.wYear if rv_time == 0 else 0,
        "sign_type": st.value,
    }}

svg_b = open({json.dumps(str(REPO_ROOT / "whitepaper/sri_yantra_whitepaper_diagram.svg"))}, "rb").read()
p7s_b = open({json.dumps(str(REPO_ROOT / "whitepaper/sri_yantra_whitepaper_diagram.svg.p7s"))}, "rb").read()
s_info_buf = ctypes.create_string_buffer(512)
rv_svg = lib.EUVerifyDataSpecific(
    ctypes.create_string_buffer(svg_b, len(svg_b)), len(svg_b),
    0,
    None,
    ctypes.create_string_buffer(p7s_b, len(p7s_b)), len(p7s_b),
    s_info_buf
)
out["svg_p7s_rv"] = rv_svg
print(json.dumps(out))
"""
        import os
        env = {**os.environ, "LD_LIBRARY_PATH": "/opt/iit/eu/sw"}
        proc = subprocess.run([sys.executable, "-c", code], env=env, capture_output=True, text=True, check=True)
        res = json.loads(proc.stdout.strip().splitlines()[-1])
        for pdf_name in ("Marma_Crypto_Whitepaper_EN.pdf", "Marma_Crypto_Whitepaper_UA.pdf"):
            r = res[pdf_name]
            self.assertEqual(r["rv_cnt"], 0)
            self.assertEqual(r["count"], 1, f"czo.gov.ua/verify EUPDFGetSignsCount != 1 for {pdf_name}")
            self.assertEqual(r["rv_ver"], 0, f"czo.gov.ua/verify EUPDFVerifyData != 0 for {pdf_name}")
            self.assertEqual(r["rv_time"], 0)
            self.assertEqual(r["bTimeAvail"], 1)
            self.assertEqual(r["bTimeStamp"], 1)
            self.assertEqual(r["year"], 2026)
            self.assertEqual(r["sign_type"], 0x10)  # CAdES-X-Long / PAdES-LTA
        self.assertEqual(res["svg_p7s_rv"], 0)

    def test_10_v13_parallel_combiner_and_256bit_bijective_inversion(self) -> None:
        """Verify v1.3.0 Parallel Multi-Hash Combiner SHA256(GOST||Kupyna||SHA256) and 256-bit bijective inversion."""
        # 1. Parallel Multi-Hash Combiner cross-verification (Python <-> Node.js dstu4145_bridge.js)
        payload_bytes = "MARMA-V1.3-PARALLEL-COMBINER-AND-256BIT-BIJECTIVE-TEST".encode("utf-8")
        sign_res = mw.call_dstu4145_bridge({
            "action": "sign",
            "secretSeed": "Test-Parallel-Combiner-Key",
            "payloadBytesHex": payload_bytes.hex(),
        })
        self.assertEqual(sign_res["status"], "ok")
        self.assertEqual(sign_res["combiner_alg"], "SHA256(GOST34311||Kupyna256||SHA256)")
        expected_sha256 = hashlib.sha256(payload_bytes).hexdigest()
        self.assertEqual(sign_res["sha256_hex"], expected_sha256)
        expected_combiner = mw.compute_parallel_combiner_hex(
            sign_res["gost34311_hex"],
            sign_res["dstu7564_hex"],
            sign_res["sha256_hex"],
        )
        self.assertEqual(sign_res["combiner_hex"], expected_combiner)

        # Verify tampering ANY single hash component in dstu4145_bridge.js fails verification
        for field in ("gost34311_hex", "dstu7564_hex", "sha256_hex", "combiner_hex"):
            tampered_pkt = dict(sign_res)
            tampered_pkt[field] = "0" * 64
            ver_res = mw.call_dstu4145_bridge({
                "action": "verify",
                "payloadBytesHex": payload_bytes.hex(),
                "dstu4145": tampered_pkt,
            })
            self.assertFalse(ver_res["dstu4145_valid"], f"Parallel combiner failed to reject tampered {field}")

        # 2. Full 256-bit bijective triangle coordinate inversion (Python <-> Browser JS across 50 digests)
        sample_digests = [hashlib.sha256(f"BIJECTIVE-256-{i}".encode("utf-8")).hexdigest() for i in range(50)]
        for idx, h in enumerate(sample_digests):
            canary_w = (idx * 19) % 1024
            for cx in (150.0, 300.0):
                tris = mw.build_perturbed_triangles_v13(h, canary_beacon_word=canary_w, cx=cx)
                rec_hex, rec_cw = mw.recover_256bit_from_triangles(tris, cx=cx)
                self.assertEqual(rec_hex, h)
                self.assertEqual(rec_cw, canary_w)
                d_coord, d_sym = mw.verify_triangle_coordinates(
                    tris,
                    [0.0] * 9,
                    cx=cx,
                    protocol="MARMA-DSTU4145-V1.3",
                    expected_hex=h,
                    canary_beacon_word=canary_w,
                )
                self.assertEqual(d_coord, 0.0)
                self.assertEqual(d_sym, 0.0)

            # Also verify 24-pair symmetric Marma mantissa encoding/decoding (264-bit capacity >= 256 bits)
            marmas = mw.encode_256bit_marma_pairs(h, canary_beacon_word=canary_w)
            self.assertEqual(len(marmas), 24)
            dec_hex, dec_c8 = mw.decode_256bit_marma_pairs(marmas)
            self.assertEqual(dec_hex, h)
            self.assertEqual(dec_c8, canary_w & 0xFF)
            for p in marmas:
                self.assertEqual(p["sym_sum_error"], 0.0)

        # Cross-verify Python v1.3 triangles against Browser JS applySignedPerturbations & recover256BitFromTriangles
        js_bijective = run_browser_js_harness(f"""
          const digests = {json.dumps(sample_digests[:10])};
          const results = digests.map((h, i) => {{
            const cw = (i * 37) % 1024;
            const tris150 = applySignedPerturbations(solveBaselineYantra(150.0), h, 150.0, "MARMA-DSTU4145-V1.3", cw);
            const inv150 = recover256BitFromTriangles(tris150, 150.0);
            const tris300 = applySignedPerturbations(solveBaselineYantra(300.0), h, 300.0, "MARMA-DSTU4145-V1.3", cw);
            const inv300 = recover256BitFromTriangles(tris300, 300.0);
            return {{
              match150: inv150.recovered_hmac_hex === h && inv150.canary_word === cw,
              match300: inv300.recovered_hmac_hex === h && inv300.canary_word === cw
            }};
          }});
          return results;
        """)
        for item in js_bijective:
            self.assertTrue(item["match150"])
            self.assertTrue(item["match300"])

    def test_11_v13_canary_leak_tracing_end_to_end(self) -> None:
        """Verify v1.3.0 Canary Leak-Tracing (--recipient-id & trace) across PDF, DOCX, SVG, and Browser JS."""
        with tempfile.TemporaryDirectory(prefix="marma_canary_") as tmp:
            tmp_dir = Path(tmp)
            target_recipient = "M&A-Party-07"
            suspect_pool = ["Investor-Alpha", "Counsel-Beta", "M&A-Party-07", "Auditor-Delta"]

            # 1. SVG with Canary recipient tag
            svg_path = tmp_dir / "confidential.svg"
            mw.sign_sri_yantra_svg(
                "CONFIDENTIAL-DEAL-MEMO-2026",
                mw.DEFAULT_SECRET,
                svg_path,
                recipient_id=target_recipient,
            )
            svg_ver = mw.verify_sri_yantra_svg(svg_path.read_text(encoding="utf-8"), recipient_id=target_recipient)
            self.assertTrue(svg_ver["verified"])
            self.assertEqual(svg_ver["protocol"], "MARMA-DSTU4145-V1.3")
            self.assertTrue(svg_ver["bijective_256bit_match"])
            self.assertTrue(svg_ver["canary_active"])
            self.assertTrue(svg_ver["canary_recipient_verified"])
            self.assertGreater(svg_ver["canary_beacon_word"], 0)

            svg_trace_keyed = mw.trace_canary_suspects(svg_path, suspect_pool, secret_key=mw.DEFAULT_SECRET)
            svg_trace_pub = mw.trace_canary_suspects(svg_path, suspect_pool, secret_key=None)
            self.assertEqual(svg_trace_keyed["matched_recipient"], target_recipient)
            self.assertEqual(svg_trace_pub["matched_recipient"], target_recipient)

            # 2. PDF with Canary recipient tag
            raw_pdf = tmp_dir / "deal_raw.pdf"
            signed_pdf = tmp_dir / "deal_canary.pdf"
            doc = pymupdf.open()
            doc.new_page().insert_text((72, 100), "Strictly Confidential M&A Term Sheet", fontsize=12)
            doc.save(str(raw_pdf))
            doc.close()

            mw.stamp_pdf_document(raw_pdf, signed_pdf, secret_key=mw.DEFAULT_SECRET, recipient_id=target_recipient)
            pdf_ver = mw.verify_pdf_document(signed_pdf, recipient_id=target_recipient)
            self.assertTrue(pdf_ver["verified"])
            self.assertEqual(pdf_ver["protocol"], "MARMA-DSTU4145-V1.3")
            self.assertTrue(pdf_ver["bijective_256bit_match"])
            self.assertTrue(pdf_ver["canary_active"])
            self.assertTrue(pdf_ver["canary_recipient_verified"])

            pdf_trace = mw.trace_canary_suspects(signed_pdf, suspect_pool, secret_key=mw.DEFAULT_SECRET)
            self.assertEqual(pdf_trace["matched_recipient"], target_recipient)
            self.assertEqual(pdf_trace["matched_count"], 1)

            # 3. DOCX with Canary recipient tag
            raw_docx = tmp_dir / "nda_raw.docx"
            signed_docx = tmp_dir / "nda_canary.docx"
            w = docx.Document()
            w.add_paragraph("Confidential NDA Copy for M&A Data Room")
            w.save(str(raw_docx))

            mw.stamp_docx_document(raw_docx, signed_docx, secret_key=mw.DEFAULT_SECRET, recipient_id=target_recipient)
            docx_ver = mw.verify_docx_document(signed_docx, recipient_id=target_recipient)
            self.assertTrue(docx_ver["verified"])
            self.assertEqual(docx_ver["protocol"], "MARMA-DSTU4145-V1.3")
            self.assertTrue(docx_ver["bijective_256bit_match"])
            self.assertTrue(docx_ver["canary_active"])
            self.assertTrue(docx_ver["canary_recipient_verified"])

            docx_trace = mw.trace_canary_suspects(signed_docx, suspect_pool, secret_key=mw.DEFAULT_SECRET)
            self.assertEqual(docx_trace["matched_recipient"], target_recipient)

            # 4. Browser JS Verifier Canary UI & 1-LSB D5.ay tamper rejection
            js_canary = run_browser_js_harness(f"""
              const buf = fs.readFileSync({json.dumps(str(signed_pdf))});
              const seal = extractMarmaSealFromBytes(new Uint8Array(buf), "deal_canary.pdf");
              renderVerificationState(seal, "deal_canary.pdf", false);
              const cleanOk = elems.statusBanner.className.includes("ok");
              const canaryStatusText = elems.valCanaryStatus.textContent;

              document.getElementById("canarySuspectsInput").value = "Investor-Alpha, M&A-Party-07, Auditor-Delta";
              runCanaryTraceUI();
              const traceResultText = elems.canaryTraceResult.textContent;

              // Shift D5.ay by a single 10-bit step (+18 nm = +0.000000018) to attempt Canary removal/spoofing
              const forged = JSON.parse(JSON.stringify(seal));
              forged.triangles[8].ay = Number((forged.triangles[8].ay + 0.000000018).toFixed(9));
              renderVerificationState(forged, "forged_canary.pdf", true);
              const forgedRejected = elems.statusBanner.className.includes("err");

              return {{ cleanOk, canaryStatusText, traceResultText, forgedRejected }};
            """)
            self.assertTrue(js_canary["cleanOk"])
            self.assertIn("ACTIVE", js_canary["canaryStatusText"])
            self.assertIn("M&A-Party-07", js_canary["traceResultText"])
            self.assertTrue(js_canary["forgedRejected"])

    def test_12_browser_sealer_to_python_cli_bidirectional_parity(self) -> None:
        """Verify that PDF and SVG files sealed in the Browser JS Studio pass Python CLI verify & trace with zero drift."""
        with tempfile.TemporaryDirectory(prefix="marma_browser_sealer_") as tmp:
            tmp_dir = Path(tmp)
            browser_pdf = tmp_dir / "browser_sealed_ma_party_07.pdf"
            browser_svg = tmp_dir / "browser_sealed_ma_party_07.svg"
            py_svg_special = tmp_dir / "py_sealed_special.svg"
            py_initial_pdf = tmp_dir / "py_initial_signed.pdf"
            browser_resealed_pdf = tmp_dir / "browser_resealed_from_py.pdf"
            py_resealed_pdf = tmp_dir / "py_resealed_from_browser.pdf"
            protocol_pdf = tmp_dir / "browser_sealed_protocol.pdf"

            secret = "DEAL-ROOM-MASTER-SECRET-2026::КЛЮЧ-108"
            recipient = "M&A-Party-07"
            suspect_pool = ["Investor-Alpha", "M&A-Party-07", "Legal-Counsel-B", "Board-Observer-X"]
            special_payload = 'КОНФІДЕНЦІЙНО · M&A <Term Sheet> & "GOLAW" \'2026\''
            special_signer = 'Баєнко Ярослав & Партнери "M&A" <Law>'

            # Prepare a Python-CLI-signed PDF first so we can also test re-sealing it in Browser Studio
            raw_pdf = tmp_dir / "raw_base.pdf"
            d_raw = pymupdf.open()
            d_raw.new_page().insert_text((72, 100), "Base Deal Document for Cross-Sealing", fontsize=12)
            d_raw.save(str(raw_pdf))
            d_raw.close()
            mw.stamp_pdf_document(raw_pdf, py_initial_pdf, secret_key="OLD-PYTHON-KEY", recipient_id="Investor-Alpha")

            # Also sign an SVG in Python CLI with XML special characters & Cyrillic
            mw.sign_sri_yantra_svg(
                special_payload,
                secret,
                py_svg_special,
                signer_key_name=special_signer,
                recipient_id=recipient,
            )

            js_seal_res = run_browser_js_harness(f"""
              document.getElementById("sealerSignerInput").value = {json.dumps(special_signer)};
              document.getElementById("sealerSecretInput").value = {json.dumps(secret)};
              document.getElementById("sealerRecipientsInput").value = "Investor-Alpha, M&A-Party-07, Legal-Counsel-B";

              // 1. Trigger 1-Click Deal-Room Leak Simulation in Browser Studio
              runOneClickDealRoomDemo();
              const demoBannerOk = elems.statusBanner.className.includes("ok");
              const demoTraceText = elems.canaryTraceResult.textContent;
              const copiesCount = generatedSealedCopies.length;

              // Save the 2nd sealed PDF copy (M&A-Party-07) to disk for Python CLI verification
              fs.writeFileSync({json.dumps(str(browser_pdf))}, Buffer.from(generatedSealedCopies[1].sealedBytes));

              // 2. Seal an SVG in the Browser Studio with XML special chars & Cyrillic for M&A-Party-07
              const svgSealed = sealSvgInBrowser({json.dumps(special_payload)}, {{
                secretKey: {json.dumps(secret)},
                signerName: {json.dumps(special_signer)},
                recipientId: {json.dumps(recipient)}
              }});
              fs.writeFileSync({json.dumps(str(browser_svg))}, Buffer.from(svgSealed.sealedBytes));

              // Verify Browser-sealed SVG in Browser Verifier
              const parsedBrowserSvg = extractMarmaSealFromBytes(svgSealed.sealedBytes, "browser_sealed.svg");
              renderVerificationState(parsedBrowserSvg, "browser_sealed.svg", false);
              const browserSvgSelfOk = elems.statusBanner.className.includes("ok");

              // Verify Python-sealed SVG (with XML special chars & Cyrillic) in Browser Verifier
              const pySvgBytes = new Uint8Array(fs.readFileSync({json.dumps(str(py_svg_special))}));
              const parsedPySvg = extractMarmaSealFromBytes(pySvgBytes, "py_sealed_special.svg");
              renderVerificationState(parsedPySvg, "py_sealed_special.svg", false);
              const pySvgInBrowserOk = elems.statusBanner.className.includes("ok");

              // 3. Re-seal the already Python-CLI-signed PDF inside Browser Studio for Board-Observer-X
              const pyPdfBytes = new Uint8Array(fs.readFileSync({json.dumps(str(py_initial_pdf))}));
              const resealed = sealPdfBytesInBrowser(pyPdfBytes, {{
                secretKey: {json.dumps(secret)},
                signerName: {json.dumps(special_signer)},
                recipientId: "Board-Observer-X"
              }});
              fs.writeFileSync({json.dumps(str(browser_resealed_pdf))}, Buffer.from(resealed.sealedBytes));
              const parsedResealed = extractMarmaSealFromBytes(resealed.sealedBytes, "browser_resealed.pdf");
              renderVerificationState(parsedResealed, "browser_resealed.pdf", false);
              const resealedInBrowserOk = elems.statusBanner.className.includes("ok");

              return {{
                demoBannerOk,
                demoTraceText,
                copiesCount,
                pdfBeacon: generatedSealedCopies[1].beaconWord,
                svgBeacon: svgSealed.sealPacket.canary.beacon_word,
                browserSvgSelfOk,
                pySvgInBrowserOk,
                parsedPySvgPayload: parsedPySvg.payload,
                parsedPySvgSigner: parsedPySvg.signer_key_name,
                resealedInBrowserOk,
                resealedSigner: parsedResealed.signer_key_name
              }};
            """)

            self.assertTrue(js_seal_res["demoBannerOk"])
            self.assertEqual(js_seal_res["copiesCount"], 3)
            self.assertIn("M&A-Party-07", js_seal_res["demoTraceText"])
            self.assertGreater(js_seal_res["pdfBeacon"], 0)
            self.assertGreater(js_seal_res["svgBeacon"], 0)
            self.assertTrue(js_seal_res["browserSvgSelfOk"])
            self.assertTrue(js_seal_res["pySvgInBrowserOk"])
            self.assertEqual(js_seal_res["parsedPySvgPayload"], special_payload)
            self.assertEqual(js_seal_res["parsedPySvgSigner"], special_signer)
            self.assertTrue(js_seal_res["resealedInBrowserOk"])
            self.assertEqual(js_seal_res["resealedSigner"], special_signer)

            # 4. Verify Browser-sealed PDF in Python CLI (both public mode and strict-key + recipient mode)
            py_pdf_pub = mw.verify_pdf_document(browser_pdf)
            self.assertTrue(py_pdf_pub["verified"], f"Python public verify failed on Browser PDF: {py_pdf_pub}")
            self.assertEqual(py_pdf_pub["protocol"], "MARMA-DSTU4145-V1.3")
            self.assertEqual(py_pdf_pub["marma_bits"], 256)
            self.assertTrue(py_pdf_pub["bijective_256bit_match"])
            self.assertTrue(py_pdf_pub["dstu4145_verified"])
            self.assertEqual(py_pdf_pub["signer_key_name"], special_signer)
            self.assertEqual(py_pdf_pub["max_coordinate_signature_drift"], 0.0)
            self.assertEqual(py_pdf_pub["max_bilateral_symmetry_error"], 0.0)

            py_pdf_strict = mw.verify_pdf_document(browser_pdf, secret_key=secret, recipient_id=recipient)
            self.assertTrue(py_pdf_strict["verified"])
            self.assertTrue(py_pdf_strict["canary_recipient_verified"])

            py_pdf_wrong_key = mw.verify_pdf_document(browser_pdf, secret_key="WRONG-KEY")
            self.assertFalse(py_pdf_wrong_key["verified"])

            # 5. Generate Official PDF Verification Protocol from Browser-sealed PDF
            proto_res = mw.generate_pdf_verification_protocol(browser_pdf, protocol_pdf, lang="en")
            self.assertTrue(proto_res["verified"])
            self.assertTrue(protocol_pdf.exists())

            # 6. Run Python CLI Canary Trace on Browser-sealed PDF
            py_pdf_trace = mw.trace_canary_suspects(browser_pdf, suspect_pool, secret_key=secret)
            self.assertTrue(py_pdf_trace["document_verified"])
            self.assertEqual(py_pdf_trace["matched_recipient"], recipient)
            self.assertEqual(py_pdf_trace["matched_count"], 1)

            # 7. Verify Browser-resealed PDF (originally signed in Python CLI) in Python CLI + Canary Trace
            py_resealed_ver = mw.verify_pdf_document(
                browser_resealed_pdf,
                secret_key=secret,
                recipient_id="Board-Observer-X",
            )
            self.assertTrue(py_resealed_ver["verified"], f"Failed on Browser-resealed PDF: {py_resealed_ver}")
            self.assertTrue(py_resealed_ver["canary_recipient_verified"])
            py_resealed_trace = mw.trace_canary_suspects(browser_resealed_pdf, suspect_pool, secret_key=secret)
            self.assertEqual(py_resealed_trace["matched_recipient"], "Board-Observer-X")

            # 8. Re-seal Browser-sealed PDF in Python CLI and verify in both Python CLI & Browser Verifier
            mw.stamp_pdf_document(browser_pdf, py_resealed_pdf, secret_key=secret, recipient_id="Legal-Counsel-B")
            self.assertTrue(mw.verify_pdf_document(py_resealed_pdf, secret_key=secret, recipient_id="Legal-Counsel-B")["verified"])
            self.assertEqual(
                mw.trace_canary_suspects(py_resealed_pdf, suspect_pool, secret_key=secret)["matched_recipient"],
                "Legal-Counsel-B",
            )

            # 9. Verify Browser-sealed SVG (with XML special chars & Cyrillic) in Python CLI + Canary Trace
            py_svg_ver = mw.verify_sri_yantra_svg(
                browser_svg.read_text(encoding="utf-8"),
                secret_key=secret,
                recipient_id=recipient,
            )
            self.assertTrue(py_svg_ver["verified"], f"Python verify failed on Browser SVG: {py_svg_ver}")
            self.assertEqual(py_svg_ver["payload"], special_payload)
            self.assertEqual(py_svg_ver["signer_key_name"], special_signer)
            self.assertTrue(py_svg_ver["bijective_256bit_match"])
            self.assertTrue(py_svg_ver["canary_recipient_verified"])
            self.assertEqual(py_svg_ver["max_coordinate_signature_drift"], 0.0)

            py_svg_trace = mw.trace_canary_suspects(browser_svg, suspect_pool, secret_key=secret)
            self.assertEqual(py_svg_trace["matched_recipient"], recipient)

            # Also verify Python-sealed SVG with XML special chars in Python CLI
            py_special_ver = mw.verify_sri_yantra_svg(
                py_svg_special.read_text(encoding="utf-8"),
                secret_key=secret,
                recipient_id=recipient,
            )
            self.assertTrue(py_special_ver["verified"])
            self.assertEqual(py_special_ver["payload"], special_payload)
            self.assertEqual(py_special_ver["signer_key_name"], special_signer)

            # 10. Tamper 1 byte in Browser-sealed PDF body -> rejected by BOTH Python CLI and Browser Verifier
            tampered_pdf = tmp_dir / "browser_sealed_tampered.pdf"
            raw_b = bytearray(browser_pdf.read_bytes())
            idx_val = raw_b.find(b"42,500,000")
            self.assertNotEqual(idx_val, -1)
            raw_b[idx_val] = ord("9")  # alter valuation from 42,500,000 to 92,500,000
            tampered_pdf.write_bytes(bytes(raw_b))

            py_tamper_ver = mw.verify_pdf_document(tampered_pdf)
            self.assertFalse(py_tamper_ver["verified"])
            self.assertIn("modified", py_tamper_ver["reason"])

            js_tamper_ver = run_browser_js_harness(f"""
              const buf = fs.readFileSync({json.dumps(str(tampered_pdf))});
              const seal = extractMarmaSealFromBytes(new Uint8Array(buf), "browser_sealed_tampered.pdf");
              renderVerificationState(seal, "browser_sealed_tampered.pdf", false);
              return {{
                tamperedTextFlag: Boolean(seal.tamperedText),
                bannerErr: elems.statusBanner.className.includes("err")
              }};
            """)
            self.assertTrue(js_tamper_ver["tamperedTextFlag"])
            self.assertTrue(js_tamper_ver["bannerErr"])


if __name__ == "__main__":
    unittest.main(verbosity=2)







